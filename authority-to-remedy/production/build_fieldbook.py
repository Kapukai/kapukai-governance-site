from pathlib import Path
import json, re, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/workspace/scratch/ceefc48769f1/rights_to_remedy_work')
OUT = ROOT / 'Kapukai_Authority_to_Remedy_Fieldbook.docx'
FONT = 'DejaVu Sans'
NAVY = '152D46'
PALE = 'EFF4F8'
GOLD = 'B3914B'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.66)
sec.bottom_margin = Inches(.64)
sec.left_margin = Inches(.72)
sec.right_margin = Inches(.72)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.28)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    s = doc.styles[name]
    s.font.name = FONT
    s.font.color.rgb = RGBColor.from_string('000000')
    s.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), FONT)
    s.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), FONT)
doc.styles['Normal'].font.size = Pt(10.7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(29)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(18)
doc.styles['Subtitle'].font.size = Pt(14)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_after = Pt(13)
doc.styles['Heading 2'].font.size = Pt(12.3)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 3'].font.size = Pt(11)
doc.styles['Heading 3'].font.bold = True
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(5)
for s in doc.styles:
    if s.type == 1:
        s.paragraph_format.widow_control = True
        ppr = s.element.find(qn('w:pPr'))
        if ppr is not None:
            for e in list(ppr):
                if e.tag == qn('w:pBdr'):
                    ppr.remove(e)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = foot.add_run('KAPUKAI GOVERNANCE LAB  |  AUTHORITY TO REMEDY  |  ')
run.font.name = FONT
run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
foot._p.append(fld)

def p(text='', boldlead=None, style=None, size=None):
    x = doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        x.add_run(boldlead).bold = True
        x.add_run(text[len(boldlead):])
    else:
        x.add_run(text)
    if size:
        for r in x.runs: r.font.size = Pt(size)
    return x

def h(text, level=2):
    return doc.add_heading(text, level=level)

def page(title, kicker=None):
    new_page = len(doc.paragraphs) > 1
    if kicker:
        x=p(kicker.upper(), size=9)
        x.paragraph_format.space_after = Pt(9)
        x.paragraph_format.page_break_before = new_page
        for r in x.runs: r.font.bold=True
    title_p=h(title,1)
    if not kicker: title_p.paragraph_format.page_break_before = new_page

def url_paragraph(url):
    pp = doc.add_paragraph()
    pp.paragraph_format.space_after = Pt(5)
    rid = doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rid)
    rr=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    ff=OxmlElement('w:rFonts'); ff.set(qn('w:ascii'),FONT);ff.set(qn('w:hAnsi'),FONT);prop.append(ff)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');prop.append(sz)
    cc=OxmlElement('w:color');cc.set(qn('w:val'),NAVY);prop.append(cc)
    rr.append(prop);tt=OxmlElement('w:t');tt.text=url;rr.append(tt);link.append(rr);pp._p.append(link)
    return pp

def table(headers, rows, widths=None, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for i,head in enumerate(headers): t.rows[0].cells[i].text=head
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),NAVY if ri==0 else (PALE if ri%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shd)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2)
                pp.paragraph_format.space_before=Pt(2)
                pp.paragraph_format.line_spacing=1.03
                for rr in pp.runs:
                    rr.font.name=FONT;rr.font.size=Pt(size);rr.font.bold=(ri==0)
                    rr.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
        if ri>0:
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    p('').paragraph_format.space_after=Pt(1)
    return t

def note(text): p(text,size=9.5)

def answer(label, lines=2):
    x=p(label);x.paragraph_format.space_after=Pt(5)
    x.runs[0].bold=True
    for _ in range(lines):
        xx=p('____________________________________________________________________________',size=9)
        xx.paragraph_format.space_after=Pt(5)

# 1
x=p('KAPUKAI GOVERNANCE LAB',size=10)
x.paragraph_format.space_before=Pt(45)
for r in x.runs:r.font.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY)
p('Authority to Remedy Fieldbook',style='Title')
p('Volume 1\nA practical method for tracing consequential public decisions',style='Subtitle')
p('A fieldbook and classroom workbook for turning one concern about public power into a traceable record, a sourced legal question, a correctly routed request, and a verifiable result.')
p('The unit of analysis is one consequential action or omission. For each unit, trace two things: the authority and conditions that permit the action, and any existing duty or remedy that may respond to an error, unlawful action, or failure to act.')
h('What you will learn')
table(['Practice','Deliverable'],[
('Record','A dated event and evidence record that distinguishes observation from inference.'),
('Classify','An action or omission tied to an actor, source, conditions, and uncertainty.'),
('Invoke documented duty','A bounded request connected to an existing duty and a valid procedural route.'),
('Seek authorized enforcement','A forum and document sequence checked against its governing rules.'),
('Verify remedy','Evidence that the authorized correction reached the decision and its consequences.')
],[1.72,5.32],10)
p('Engineering led  |  AI assisted  |  Human directed',size=10)
p('First classroom module  •  Edition dated 9 October 2026',size=9.5)
note('Educational method and fictional exercises. The cited source controls over a summary. A real case requires current, jurisdiction-specific verification of authority, procedure, deadlines, and available relief.')

# 2
page('How to use this fieldbook')
p('Start with a single decision that changes a person’s position: a restriction, an adverse determination, a denial, a missing service, or a required decision that has not occurred. Describe the effect in ordinary words before choosing a legal label. A complete legal strategy can involve several linked cards; this volume teaches how to make one card reliable.')
h('The reading and practice route')
table(['Pages','Use'],[
('3 to 12','Read the ten short chapters before class. They explain the two traces, lawful documentation, duty analysis, forum selection, and remedy verification.'),
('13 to 14','Study a real-law conditional duty example and a map of required government records in Florida dependency proceedings.'),
('15 to 16','Complete the two-page Action Duty Remedy Card. Keep unknowns visible.'),
('17 to 18','Use the source, deadline, and atlas registers to make the card maintainable.'),
('19 to 21','Work the three fictional exercises. The facts and training rules are supplied; do not import private case details.'),
('22','Use the facilitator key to assess reasoning and identify missing information.'),
('23 to 25','Consult the primary-source register. Open the current source before applying it to an actual situation.')
],[.75,6.29],10)
h('The relationship to the Living Blueprint')
p('Kapukai’s Living Blueprint organizes a wider system through domain, scope, vision, stakeholders, priorities, heuristics, functions, components, measures, coupling, coherence, tradeoffs, and learning. This fieldbook applies that architecture to one traceable exercise of public power. It does not replace the larger blueprint or supply a complete national court atlas.')
p('For this class, the design priorities are safety and rights, lawful authority, reliable evidence, effective participation, timely remedy, and then efficiency. These priorities guide the learning environment. They do not override legal standards, judicial orders, or a court’s assigned jurisdiction.')
h('Working conventions')
p('Use “verified,” “disputed,” “unknown,” “not applicable,” and “proposed” deliberately. Preserve the difference between a statute, regulation, court rule, order, agency policy, and an engineering recommendation. A request may invoke an existing duty; it does not create new public authority by assertion.')
note('The printable PDF has writing space. The Word edition is editable. Neither format is a court form, and the fictional exercises contain no real filing deadlines.')

# 3
page('One decision and two traces','Chapter 1')
p('A broad complaint can contain several different legal problems. Narrowing the unit makes them testable. “The agency harmed our family” becomes “On this date, this actor imposed this contact restriction using this document.” A separate omission card might ask whether an identified actor failed to provide a review that a specific source required.')
table(['Permission to act','Duty and remedy'],[
('Who is acting and in what role?','Who owes the asserted duty and to whom?'),
('What source grants this specific power?','What source requires this act or permits this review?'),
('What facts and procedures must exist first?','What event triggers the duty, and is it satisfied?'),
('What scope, duration, and limits apply?','What deadline, discretion, or exception applies?'),
('What evidence supports each condition?','What authorized forum can address a failure?'),
('What action actually occurred?','What relief can that forum actually order?')
],[3.52,3.52],10)
p('The two traces meet at the same evidence record. A missing predicate may undermine an action. A missing response may support a procedural request. Neither conclusion follows merely from dissatisfaction with the result. The relevant source and facts must support the particular claim.')
h('Define the smallest useful unit')
p('Assign a stable decision identifier, the affected person or protected pseudonym, the actor, the action or omission, the date, the document or system output, and the practical effect. Link related units instead of combining them into a single undifferentiated allegation.')
p('For an omission, describe the act that should have occurred. “Failure to protect” is broad. “Failure to issue the written determination required by section X after a complete request was received” is narrower. The second description can be tested for scope, trigger, timing, and remedy.')
h('A stopping rule that preserves accuracy')
p('If the source, actor, deadline, or route cannot be verified, record the gap and identify the next verification step. Do not turn an unknown into a finding of misconduct. Do not treat an unknown as permission to ignore a binding order. The useful output may be a precise question for counsel or the responsible office.')
note('Workbook output: one identified decision and two incomplete but explicit traces. Relevant legal distinctions appear in Chapters 2, 3, 7, and 8.')

# 4
page('Trace public power to the assigned office','Chapter 2')
p('The Declaration of Independence explains a political justification for government grounded in the people and their rights. It is a foundational statement, not a substitute for the operative source that grants an official power or creates a modern procedural remedy. Begin there for purpose; continue to the constitutions and laws for the enforceable structure. [S1]')
table(['Layer','What to locate','What it does not establish by itself'],[
('The people and founding purpose','The Declaration and constitutional text.','The specific authority of an agency employee or a ready-made cause of action.'),
('Federal constitutional structure','Enumerated federal powers, separation of powers, rights, supremacy, and reserved powers.','That every local child welfare decision is a federal agency action.'),
('State constitutional structure','The state’s branches, courts, rights, and allocation of powers.','Unlimited authority to act contrary to controlling federal law.'),
('Lawful delegation','Applicable statute, valid regulation, local enactment, or authorized order.','That every action of a legitimate institution is within its powers.'),
('Assigned office and person','Appointment, role, territorial reach, subject matter, and decision-specific delegation.','That a recommendation is an order or that a contractor has every power of its public client.'),
('Specific action','Required findings, conditions, process, limits, and current evidence.','That a welfare objective satisfies all legal prerequisites.')
],[1.37,2.9,2.77],9.6)
p('Federal and state authority are parallel constitutional systems with defined relationships. The Supremacy Clause addresses controlling federal law; the Tenth Amendment recognizes reserved powers. The Fourteenth Amendment constrains state action through due process and equal protection. Florida Article I section 1 locates political power in the people; sections 9 and 21 address due process and access to courts. [S2, S3]')
h('Make every connection explicit')
p('For each connection, write “this source authorizes this actor to do this act when these conditions exist.” Preserve limits and exceptions. A general mission statement explains purpose; it may not confer the coercive power being examined.')
note('The blueprint is a tracing method, not a claim that a person can declare themselves outside lawful jurisdiction. Constitutional questions require the procedural vehicle and forum that can lawfully decide them.')

# 5
page('Turn a legal provision into a testable rule','Chapter 3')
p('Read the operative text together with its definitions, exceptions, cross-references, effective date, and controlling interpretations. A sentence containing “shall” may still depend on eligibility, a triggering event, or a prescribed procedure. A provision allowing an actor to choose does not necessarily allow them to ignore a separate duty to consider or decide.')
table(['Rule element','Question to answer','Record on the card'],[
('Actor','Exactly who is addressed?','Office, officer, tribunal, provider, or defined class.'),
('Required or permitted act','What verb describes the legal act?','Decide, serve, hear, disclose, review, provide, correct, or another exact act.'),
('Trigger','What must happen first?','A filing, receipt, finding, request, event, status, or deadline.'),
('Predicates','What conditions must be established?','Each element separately, with supporting and contrary evidence.'),
('Timing','When must the act occur?','Exact time rule and the event that starts the clock.'),
('Choice and limits','What discretion or exceptions remain?','Discretion over process, timing, outcome, or none, as supported.'),
('Consequence and route','What follows if the duty is unmet?','Authorized request, review, enforcement, or an unresolved remedy question.')
],[1.16,2.6,3.28],9.6)
h('Three separate conclusions')
p('First, determine whether a legal duty exists. Second, determine whether its conditions are met. Third, determine whether this person can seek this remedy in this forum. A legal duty does not automatically create a private damages claim, a right to compel prosecution, or jurisdiction in any court a person selects.')
h('Write a bounded predicate statement')
p('Use a form such as: “If A receives a complete request of type B from a qualifying person, and exception C does not apply, source D requires act E within time F.” Then connect every factual term to a dated item. Where the source uses a legal standard rather than a fixed formula, preserve that judgment rather than replacing it with a homemade score.')
note('The federal APA example in Chapter 8 shows why “required to decide” and “required to decide in my favor” are different propositions. The teaching method does not convert all law into a mechanical algorithm.')

# 6
page('Create a reliable and lawful record','Chapter 4')
p('A record should enable another person to reconstruct what happened without accepting the recorder’s conclusion on trust. Make contemporaneous notes when possible. Capture the date, local time and time zone, location or channel, participants and roles, the observable act, words remembered, documents received, and immediate practical effects.')
p('Use quotation marks only for an exact quotation. Label recollection, paraphrase, inference, and information learned from someone else. Preserve an original file and work from a copy. A hash can help show that a file has not changed; it does not prove that the file’s contents are true.')
table(['Record type','Useful practice','Limit to preserve'],[
('Notes','Date the entry and identify what was directly observed.','A later recollection is not a contemporaneous record.'),
('Documents','Keep complete versions, envelopes, notices, attachments, and delivery records.','A partial screenshot may omit conditions or context.'),
('Audio or video','Verify consent, location, communication, and court rules before recording.','The technical ability to record does not establish permission.'),
('System records','Preserve available timestamps, recipients, status changes, and audit history.','A database entry may show that a statement was entered, not that it is true.'),
('Evidence index','Assign IDs and summarize relevance without changing originals.','Repeated copies from one source are not independent corroboration.')
],[1.04,3.0,3.0],9.7)
h('Florida recording practice')
p('Florida section 934.03 regulates interception of covered communications. One statutory route permits interception when all parties have given prior consent; definitions and other exceptions require fact-specific analysis. A public-interest purpose is not blanket permission to record secretly. Courtroom and proceeding-specific rules are a separate question. For this class, use notes, lawfully obtained records and transcripts, and explicitly consensual simulations. [S16]')
h('Protect the record and the people')
p('Keep child identifiers, protected addresses, medical details, and case records out of public class materials. Use fictional cases and pseudonyms. Limit access and distinguish a protected working copy from a version authorized for sharing. Florida dependency records have statutory confidentiality rules; “evidence” does not mean “publishable.” [S13]')
note('Workbook output: an event log and evidence inventory with source, date, access limits, and evidence status. Ask for the current local recording or transcript procedure through the proper channel when uncertain.')

# 7
page('Classify conduct without outrunning the evidence','Chapter 5')
p('Classification should help select the next question. It should not turn a suspicion into a finding. Start with conduct: an order entered, a notice not received, a database field changed, a hearing not held, or a service not made available. Then classify the possible defect and identify the source needed to test it.')
table(['Possible category','Question to test','Evidence often needed'],[
('Identity or record error','Is this the right person, relationship, event, or version?','Source record, matching identifiers, amendment history.'),
('Authority or scope','Could this actor take this specific action?','Delegation, current order, role, jurisdictional facts.'),
('Missing factual predicate','Was a required condition supported?','Element-by-element evidence and contrary material.'),
('Process failure','Was a required notice, hearing, reason, or opportunity supplied?','Rule, service record, docket, decision, access history.'),
('Required act not done','Did a triggered mandatory duty remain unperformed?','Exact source, request, receipt, timing, actor response.'),
('Discretion challenged','Was a permitted choice exercised within lawful limits?','Reasoning, record, governing standard of review.'),
('Implementation failure','Did the operative decision reach actual practice?','Order, instructions, system state, delivery and access tests.')
],[1.32,2.78,2.94],9.8)
h('Keep fault and remedy separate')
p('The same record error can arise from a typo, a badly designed process, negligent verification, or deliberate falsification. You may be able to request correction before establishing intent. A criminal accusation, disciplinary complaint, civil claim, and motion in an existing case have different elements, decision-makers, and effects. One does not automatically substitute for another.')
h('Classify omissions precisely')
p('Identify the missed act and the responsible actor. An unavailable service may concern a provider’s performance, an agency’s provision responsibility, a parent’s participation requirement, or several of these. Do not charge the wrong actor with a duty they do not have. Do not count a barrier controlled by one participant as evidence of another participant’s unwillingness without analyzing the facts.')
note('Workbook output: a provisional category, the evidence supporting it, evidence against it, and an explicit uncertainty. Use ordinary descriptions until the applicable legal category is verified.')

# 8
page('Invoke an existing duty with a bounded request','Chapter 6')
p('The user-facing shorthand “obligate” means invoking a duty the law or an operative order already imposes. It does not mean that sending a declaration, affidavit, invoice, or demand creates a government obligation. An affidavit can present facts; it does not adjudicate them or compel agreement through silence.')
h('The request as a traceable package')
table(['Component','What it contains'],[
('Who and what','The requestor’s relevant status, the specific action or omission, and the official or office addressed.'),
('Governing source','The exact provision or order, version, and paragraph, with the trigger and required act.'),
('Facts and exhibits','A short chronology; evidence IDs; relevant contrary facts; clear statements of uncertainty.'),
('Requested action','The act the recipient has authority to perform, phrased precisely enough to verify.'),
('Procedure','Required form, filing destination, fee or waiver process, service, signature, and supporting material.'),
('Timing and preservation','The applicable deadline; proof of receipt or service; separately tracked review and appeal deadlines.'),
('Response and next route','Where to send a response; the source-based review route if the issue remains unresolved.')
],[1.34,5.7],10)
p('A useful request might seek correction of a named field, issuance of a required written determination, provision of an accessible service, clarification of an order, or consideration of a specified form of relief. State why the actor is the right recipient. If the duty is disputed, ask the appropriate question rather than declaring the conclusion settled.')
h('Do not assume that one clock pauses another')
p('A complaint, records request, reconsideration request, or correction request may have no effect on an appeal deadline. Florida Juvenile Rule 8.265(b)(3) expressly states that rehearing does not toll appeal time. Rule 8.265(b)(4) also addresses preserving a failure to make required findings in a final order. Check the actual order and current rules; do not import ordinary civil rehearing assumptions. [S11, S12]')
note('Workbook output: a request outline, a verified procedural checklist, and proof of delivery or service. The fact that a clerk accepts a filing does not prove that the forum can grant the requested relief.')

# 9
page('Build an atlas of forums and routes','Chapter 7')
p('A usable atlas is more than a map of buildings. It connects a problem, governing legal system, decision-maker, procedural vehicle, deadline, service rule, standard of review, and available relief. Geography matters, but subject matter, procedural stage, and the identity of the actor matter too.')
table(['Route','What to verify before using it'],[
('Existing court case','Court and division, current orders, party status, authorized motion or application, notice and service, hearing procedure.'),
('Administrative process','The agency, benefit or regulatory program, internal review mechanism, exhaustion requirements, and record rules.'),
('Appellate review','Reviewable order, court with appellate jurisdiction, filing vehicle, rendition and deadline rules, record, and any stay procedure.'),
('Extraordinary relief','The specific writ, prerequisites, correct respondent and court, adequacy of other remedies, and limits on discretion.'),
('Independent civil claim','Cause of action, standing, jurisdiction, defendant, immunity, limitations, service, and available remedy.'),
('Oversight or reporting','The office’s actual remit and what a complaint can accomplish. Track separately from judicial review deadlines.')
],[1.45,5.59],10)
h('A Florida dependency example of route separation')
p('Dependency proceedings belong to Florida’s state court structure under chapter 39 and the relevant court rules. Chapter 39 contains distinct provisions on custody, shelter proceedings, adjudication, disposition, review, and appeal. The applicable path depends on the decision and procedural stage. The federal APA does not govern a Florida dependency judge merely because public power is involved. [S8-S12]')
p('For Hillsborough County, the atlas identifies the Thirteenth Judicial Circuit and the relevant dependency division, then verifies the actual case, administrative orders, filing channel, and local procedures. The Second District Court of Appeal includes Hillsborough in its geographic jurisdiction; the particular order and route still determine whether and how review is available. [S14, S15]')
note('Workbook output: a route record with a primary-source link and verification date for every procedural claim. An atlas entry may remain “unverified” until each necessary field is checked.')

# 10
page('Seek enforcement within the remedy that exists','Chapter 8')
p('A duty is only one part of an enforcement analysis. Ask whether this claimant may invoke the remedy, which forum has jurisdiction, what process is required, what standard applies, and what relief the decision-maker can grant. The appropriate result may be a new decision, a hearing, a correction, an injunction, a remand, or other authorized relief. It is not always damages or the desired substantive outcome.')
h('A bounded federal APA example')
p('For covered federal agency matters, 5 U.S.C. 555(b) addresses concluding a matter presented to an agency within a reasonable time. Section 702 provides a review framework for persons suffering legal wrong or adversely affected or aggrieved within a relevant statute, with important conditions. Section 704 addresses reviewable agency action, including final agency action for which there is no other adequate remedy in a court. Section 701 identifies exceptions. Read them together. [S4-S6]')
p('Section 706 distinguishes compelling agency action unlawfully withheld or unreasonably delayed from reviewing completed agency action on specified grounds. In Norton v. Southern Utah Wilderness Alliance, the Supreme Court limited section 706(1) to discrete agency action that the agency is legally required to take. Broad dissatisfaction with administration is not that same claim. [S7]')
table(['If the asserted defect is','The analysis focuses on'],[
('A missing required decision','The discrete required act, source, trigger, agency responsibility, timing, and available review.'),
('An unfavorable completed decision','Reviewability, finality, applicable review vehicle, legal grounds, record, and authorized relief.'),
('A decision committed to discretion','The limits of review and whether any separately required procedure or act is identifiable.'),
('A state court order','The applicable state court and review system, not an assumed federal APA complaint.')
],[2.25,4.79],10)
h('Request what the reviewing body can actually do')
p('If the duty is to issue a determination, an order to decide may leave the substantive choice to the agency. If the issue is a defective court order, a regulator or complaint office may lack power to set it aside. Match the requested remedy to the forum before investing in a document sequence.')
note('Workbook output: a remedy hypothesis with its legal source and eligibility limits. This chapter is a federal example, not a universal filing recipe.')

# 11
page('Verify that remedy reaches the consequence','Chapter 9')
p('A response is not necessarily a remedy. A corrected source record is not necessarily a corrected downstream decision. Define success before sending the request: which field, order, access right, service, or action should change, who controls it, and what evidence would establish completion?')
table(['Remedy layer','Evidence of completion','Common unresolved remainder'],[
('Source','The inaccurate field or status is corrected or properly marked.','The old information is still the default view.'),
('Decision','The authorized decision-maker reconsiders or changes the affected decision as required.','The corrected fact has not been evaluated for its legal effect.'),
('Recipients','Known recipients receive the correction where authorized or required.','Copied allegations remain active elsewhere.'),
('Implementation','The operative order or decision changes actual practice.','Staff or contractors continue an earlier instruction.'),
('Human access','The affected person can use the access or support that has been restored.','A favorable document produces no practical change.'),
('Closure and learning','Unresolved effects are assigned an owner and review date.','The same input can trigger the same mistake again.')
],[1.25,2.90,2.89],9.8)
h('Preserve limits on propagation')
p('“Correction should travel as far as the error” is a design objective. Whether an organization must notify particular recipients, amend records, or reconsider a decision depends on applicable law, orders, contracts, and policy. Label that distinction on the card. Do not promise that a correction automatically erases lawful records or changes every independent decision.')
h('Measure what matters')
p('Useful measures include time to verified response, unresolved predicate count, records with source provenance, affected recipients reached, implementation completion, and reoccurrence of the same error. A low complaint count can reflect inaccessible reporting. Fast closure can conceal an incomplete remedy. No single count proves child well-being or fairness.')
p('Retest with the least intrusive evidence needed. For access restoration, verify the access. For a corrected record, inspect the permitted current view. For a service, confirm usable availability. Stop claiming completion when any required layer remains unverified.')
note('Workbook output: a remedy acceptance test, an implementation owner, a review date, and the remaining uncertainties.')

# 12
page('Teach the method through preparation and practice','Chapter 10')
p('The classroom should turn the method into a repeatable skill. Assign the short lesson and one source-reading exercise before the meeting. Use class time to compare interpretations, expose missing facts, and rehearse lawful documentation, bounded requests, and remedy tests. All practice cases in this volume are fictional.')
table(['Stage','Participant task','Evidence of learning'],[
('Before class','Watch or read the lesson; complete Case A; use Chapters 1 to 8 for deeper reading.','A short completed trace, two deadline lines, and one route question.'),
('Individual opening','Record five supplied facts without adding conclusions.','Observation, source, and inference remain distinct.'),
('Small groups','Assign fact recorder, authority researcher, procedural navigator, independent reviewer, and affected-person perspective or timekeeper.','A completed Action Duty Remedy Card with cited training rules.'),
('Role rotation','Swap roles and examine another group’s card.','The reviewer can reconstruct the claim and identify unsupported steps.'),
('Simulation','Deliver a bounded request; respond using the supplied rules; verify a mock correction.','A complete request, response record, and remedy acceptance test.'),
('After class','Revise the card from feedback and complete a short reflection.','The learner explains what changed and what remains unverified.')
],[1.05,3.0,2.99],9.7)
h('Roles for producing the class')
p('A source researcher verifies legal text and updates. An instructional designer creates scenarios and assessment. A writer prepares the fieldbook and narration. A slide designer makes the rule relationships readable. A media editor handles voice, captions, pacing, and licensed music. An independent reviewer checks legal scope, factual fidelity, accessibility, and the rendered outputs. A human owner approves publication and future changes.')
h('Keep the lesson exact and the practice active')
p('Use native slide text and diagrams for statutes, predicates, and court routes. Video imagery may introduce a scenario, but it should not depict generated text as a genuine court record. Keep narration clean and music quiet; captions and a transcript must preserve every material qualification. The source register is the shared reference for the website, fieldbook, slides, and video.')
note('Suggested class length: 90 minutes. Use 8 minutes for retrieval, 10 for demonstration, 20 for team work, 12 for cross-review, 17 for new evidence and role rotation, 12 for simulation, 6 for remedy checks, and 5 for individual exit proof. Keep private child and family material out of class.')

# 13
page('Trace a conditional duty in Florida law','Worked example')
p('Florida section 39.701(2)(d)3 provides a narrow example of a trigger and different resulting powers. At the relevant dependency judicial review, the court assesses whether the social service agency complied with obligations in the written case plan. The controlling predicate is the court’s assessment, not a private allegation. [S10]')
table(['Trace field','Educational entry'],[
('Unit','The agency allegedly did not perform a specified written case-plan obligation. Analyze each obligation separately.'),
('Evidence to organize','Operative case-plan provision; assigned actor; provider availability; communications; efforts made; any contrary records and later orders.'),
('Forum and vehicle','Existing circuit dependency proceeding. Verify party status, hearing posture, Rule 8.415, local procedure, notice, service, and evidentiary requirements. [S8, S11]'),
('Legal trigger','The court concludes that the agency has not complied with the obligations specified in the written case plan. [S10]'),
('Discretionary branch','The court may find the agency in contempt. A finding of noncompliance does not make contempt automatic. [S10]'),
('Required branches','The court must order the agency to submit plans for compliance and require it to show why the child could not safely return home. [S10]'),
('Required court record','Rule 8.415(f)(7) calls for a written review order with the relevant findings, future course, and next hearing. Rule 8.260 governs written orders and transmission. [S11]'),
('Implementation evidence','Obtain the entered order; identify the agency’s ordered response; track actual provision of the required service and any later court determination.'),
('Scope limit','Return-home relief has its own safety conditions. Agency noncompliance alone does not establish all conditions for return. [S10]')
],[1.52,5.52],9.6)
p('A learner’s proposed document sequence is: assemble the operative plan and evidence; use the authorized request and service procedure; seek the relevant finding; obtain the written order; verify the ordered outputs and their implementation. These organizational steps are a proposed workflow. The conditional legal duties come from the cited sources.',size=10)
note('Track review and appeal separately. Do not treat the submission of a request, a clerk’s receipt, or the scheduling of a hearing as proof that relief has been granted.')

# 14
page('Locate the required government record','Florida dependency checkpoint map')
p('Each row starts with its own legal trigger. Some branches end, repeat, or change direction. This is a checkpoint map for reading a docket, not a claim that every case must pass through every row or that an error produces an automatic outcome.')
table(['Trigger or stage','Record or action to locate','Primary source'],[
('Officer takes a child into custody','Facts supporting removal and the officer’s full written report to the department.','39.401(2) [S8]'),
('Department receives or takes custody','Attorney review of shelter probable cause; appropriate return or shelter petition and hearing branch.','39.401(3) [S8]'),
('Continued shelter sought','Shelter petition, notice and hearing, order containing required findings, and next review notice.','39.402(5) to (8), (16) [S9]'),
('Dependency adjudication sought','Written sworn petition identifying the relevant acts or omissions and persons.','39.501(1), (3) [S8]'),
('Contested allegations decided','Dismissal or dependency findings and order, subject to the applicable statutory branches.','39.507(4) to (6) [S8]'),
('Disposition occurs','Filed and served case plan and assessment; written disposition with placement, conditions, monitoring, and next review.','39.521(1) [S8]'),
('Judicial review occurs','Social study and required reports, notice and service, hearing record, written review order.','39.701(1), (2); Rule 8.415 [S10, S11]'),
('Court finds agency case-plan noncompliance','Court orders agency compliance plans and requires an explanation concerning safe return; contempt remains discretionary.','39.701(2)(d)3 [S10]'),
('Entered order is challenged','The authorized rehearing, relief, appeal, or writ route with its own preservation, deadline, record, and service requirements.','39.510; Rules 8.265, 8.270, 9.146 and applicable review rules [S11, S12]')
],[1.54,3.72,1.78],9.35)
h('Do not infer the next document from a slogan')
p('“Best interests” does not tell you which stage is underway, who bears a duty, or what document must follow. The source does. Record the exact subsection, required contents, recipients, and timing for the actual branch. A treatment request, change in placement, and judicial review may require separate cards even within one case.',size=10)
note('The source register links to current official materials reviewed for this edition. A real calendar requires the actual event dates, governing effective law, court rules, and applicable exceptions.')

# 14
page('Action Duty Remedy Card','Reusable worksheet 1 of 2')
p('Complete one card for one consequential action or omission. Use evidence IDs and source IDs. “Unknown” is a valid entry when it is paired with the next verification step.',size=10)
table(['Identification','Entry'],[
('Decision ID and linked cards','\n'),('Prepared by and date','\n'),('Affected person or protected pseudonym','\n'),('Actor and exact role','\n'),('Action or omission and date','\n'),('Practical effect on the person','\n')
],[2.0,5.04],10)
table(['Permission trace','Entry'],[
('Source of power and exact section','\n'),('Jurisdiction and assigned role','\n'),('Required predicates and procedures','\n\n'),('Supporting evidence IDs','\n'),('Contrary evidence and uncertainties','\n'),('Scope duration and limits','\n')
],[2.0,5.04],10)
note('Keep evidence status and legal-source type explicit. Use the conventions on page 2.')

# 15
page('Action Duty Remedy Card','Reusable worksheet 2 of 2')
table(['Duty and remedy trace','Entry'],[
('Duty source and exact required act','\n'),('Responsible actor and beneficiary','\n'),('Trigger and proof it occurred','\n'),('Deadline and remaining discretion','\n'),('Exceptions or contrary evidence','\n'),('Requested action and legal basis','\n'),('Forum vehicle and eligibility','\n'),('Required attachments and service','\n'),('Independent clocks and tolling source','\n'),('Proof of filing receipt or service','\n'),('Remedy acceptance test','\n'),('Downstream recipients or decisions','\n'),('Implementation owner and review date','\n')
],[2.0,5.04],10)
p('Reviewer: __________________________  Date: __________________________',size=10)
p('Next verified action: _______________________________________________________',size=10)
note('This card records reasoning and verification. It is not a finding of illegality or a substitute for procedure.')

# 16
page('Keep sources and deadlines together','Reusable registers')
p('A source register lets another person verify a rule. A deadline register prevents an unanswered request from quietly consuming time for another procedure. Use a separate row for every clock and record the rule that starts and changes it.')
h('Source register')
table(['ID and type','Authority and pinpoint','Version and checked date','Claim supported'],[
('S___\nStatute rule order policy','\n\n','\n\n','\n\n'),
('S___\nStatute rule order policy','\n\n','\n\n','\n\n'),
('S___\nStatute rule order policy','\n\n','\n\n','\n\n')
],[1.07,2.21,1.67,2.09],9.5)
p('For each source, also retain its full official URL or docket locator, jurisdiction, effective date, relevant definition or exception, and the person responsible for rechecking it.',size=10)
h('Deadline register')
table(['Event and clock','Source and start event','Calculation and uncertainty','Owner and proof'],[
('D___\nFiling or response','\n\n','\n\n','\n\n'),
('D___\nReview or appeal','\n\n','\n\n','\n\n'),
('D___\nService or hearing','\n\n','\n\n','\n\n')
],[1.25,2.04,2.07,1.68],9.5)
h('Calculation checks')
p('Verify the start event, calendar or business-day rule, holidays, time zone, filing method, applicable service additions, rendition, and tolling. Have an important real deadline independently checked. Record an unresolved rule question instead of an unlabeled guessed date.')

# 17
page('Court and agency atlas entry','Reusable route schema')
p('Create one entry for one procedural route, not merely one courthouse. A single institution may need several entries because motions, appeals, records access, and emergency applications have different requirements.')
table(['Atlas field','What the entry must hold'],[
('Route ID and problem','Unique ID; decision or omission type; linked Action Duty Remedy Card.'),
('Legal system and location','Federal state tribal or other system; territory; county or district; relevant forum relationships.'),
('Actor and institution','Agency or court; division; assigned official or case; verified contact and filing source.'),
('Jurisdiction and eligibility','Subject matter; person permitted to invoke the route; necessary party status; respondent or defendant.'),
('Source of duty and remedy','Exact provisions and decisions; distinguishing the asserted duty from the source of review power.'),
('Procedural vehicle','Application motion petition notice complaint or other authorized document; current approved form if one exists.'),
('Prerequisites','Exhaustion finality notice demand authorization or other conditions; exceptions separately sourced.'),
('Clock and service','Trigger; deadline; calculation rule; service recipients; method; proof; verified tolling effects.'),
('Record and hearing','Attachments; evidentiary requirements; record designation; hearing request and scheduling rules.'),
('Available relief','What the forum can order; limits; separate interim or stay request if permitted and needed.'),
('Review path','Reviewable decision; next forum; next document; next clock.'),
('Implementation','Responsible office; communication to recipients; evidence that relief reaches actual practice.'),
('Maintenance','Official links; last checked date; source owner; confidence; update trigger; unresolved fields.')
],[1.6,5.44],9.7)
note('Map colors should represent verification status, not implied authority: verified current route; conditional route; unresolved route; proposed feature. A displayed line between forums must identify the legal relationship it represents.')

# 18 onward generated from curriculum data or local core draft below
CASES = [
    {
        'title':'Exercise A Correct a copied identity error',
        'facts':'It is training day D8. At D0, fictional Riverton Family Services changed transport booking priority for Alex Reed’s account R-184 using three missed-pickup entries. The original entries identify a different account, R-814. An import changed the number; three later summaries repeat that import. Alex filed a documented correction request on D6 and received receipt C-206. It did not expressly request review of the transport decision. At D7 the transport contractor still held the flag.',
        'rules':'TR-A1: a written accuracy determination is due 10 training days after a sufficiently identified correction request. TR-A2: internal review must be requested by D15; correction does not pause that clock. TR-A3: the reviewer may suspend the priority change. TR-A4: a confirmed error requires source correction and notice to known recipients within 3 training days. No court route is supplied. All rules and clocks are fictional.',
        'evidence':'A1 adverse notice; A2 original log; A3 imported summary; A4 correction request and receipt; A5 contractor acknowledgment; A6 training rules.',
        'questions':['Which evidence supports an identity error, and which reports are derivative? Is the correction response overdue on D8?','Write both due dates and separate the requested correction from the decision review. Is interim suspension guaranteed?','Define source, recipient, and operational remedy tests. What authority is missing before proposing a court filing?']
    },
    {
        'title':'Exercise B Separate service access from participation',
        'facts':'It is training day D12. A fictional court’s D0 order requires Jordan Lane to contact the agency by D3 and attend an available approved course by D30. The agency must make an approved course available without participant cost by D7. Jordan contacted the acknowledged caseworker on D2. On D5 the only provider reported no approved place before D45; an offered paid class was not approved. A D9 case entry says Jordan failed to engage. Review is scheduled for D21.',
        'rules':'TR-B1: a request to clarify, modify, or enforce the order goes to the assigned court in the existing case, with supporting exhibits and delivery to listed parties under the exercise service instructions. Clerk receipt is not relief. TR-B2: the court decides agency-compliance and deadline requests. TR-B3: a separate agency correction route exists but no response deadline is supplied. All rules and clocks are fictional.',
        'evidence':'B1 order; B2 acknowledged D2 contact; B3 provider response; B4 case entry; B5 procedural card.',
        'questions':['Separate Jordan’s duties from the agency’s duties. Which have been met, appear unmet, or remain conditional?','Describe the record correction request and the court request. What continues while either is pending?','Define a usable-service acceptance test. Identify missing facts and the documents needed to change the operative instructions.']
    },
    {
        'title':'Exercise C Obtain a required determination',
        'facts':'It is training day D24. The invented Federal Family Records Review Office acknowledged Taylor Brooks’s complete status-review application on D0 and said a written determination was due D20. Its portal still shows pending; the supplied correspondence contains no determination. It is unknown whether a decision exists but was not delivered. An unverified online template claims that silence guarantees approval and creates $10,000 per day in damages.',
        'rules':'TR-C1: after a complete application, the Office must issue a written determination within 20 training days. It may grant status if the eligibility criteria are met. The packet does not establish any special review statute, adequate alternative remedy, exhaustion requirement, judicial route, service rule, extension, or exception. These are fictional rules and clocks, not an actual APA deadline.',
        'evidence':'C1 training statute; C2 completeness acknowledgment; C3 pending status; C4 unverified template; C5 missing-information card.',
        'questions':['Identify the discrete potentially overdue act and its trigger. What facts must still be verified?','Write a bounded request. Explain why the template’s consequences do not follow from the supplied rule.','List the research gates for an actual federal review route. Can an unfavorable determination satisfy this narrow process remedy?']
    }
]
for idx,c in enumerate(CASES):
    page(c['title'],'Fictional classroom worksheet '+str(idx+1))
    p(c['facts'])
    p(c['rules'],boldlead=None,size=10)
    p('Evidence packet: '+c['evidence'],boldlead='Evidence packet:',size=10)
    p('Roles: fact recorder; authority researcher; procedural navigator; independent reviewer; affected-person perspective and timekeeper.',size=9.3)
    for i,q in enumerate(c['questions']):answer(str(i+1)+'. '+q,2)
    note('Do not use personal case details. Cite the supplied training rule for an answer, and mark every real-law question as requiring current verification.')

# 21
page('Facilitator key and assessment')
h('Exercise A')
p('Expected trace: R-184 and R-814 are different accounts. Repeated summaries are derivative, not three independent sources. The D6 correction request makes the TR-A1 determination due D16, so it is not overdue at D8. Internal review is due D15 and is not paused. Preserve it separately. Interim suspension is discretionary. If error is confirmed, TR-A4 requires source correction and notice to known recipients. Verify the transport decision and contractor’s operative flag separately. Judicial routing remains unknown.')
h('Exercise B')
p('Expected trace: Jordan met the contact requirement on D2, before D3. The agency’s availability duty was due D7 and appears unmet; check for omitted referrals or a later order. Attendance depends on an available approved course. The D45 slot cannot satisfy the current D30 requirement without further action. Seek factual correction, actual access, and an authorized court decision on compliance or deadline modification. Keep the D21 review and comply with unaffected duties. An agency email cannot modify the court order.')
h('Exercise C')
p('Expected trace: the written determination appears four training days late. Verify whether it exists, whether delivery failed, and whether an applicable exception changes the analysis. No rule grants approval or automatic damages through silence. A duty to decide is distinct from eligibility for the desired status. An actual APA route requires coverage, standing, jurisdiction, reviewability, special-law and alternative-remedy analysis, and proper procedure. A lawful adverse determination can complete this narrow process remedy while leaving a separate substantive challenge. [S4-S7]')
table(['Assessment dimension','Strong evidence of learning'],[
('Factual discipline','Separates observed facts, reports, inferences, and unknowns.'),
('Authority and duty','Names the actor, source, trigger, required act, and remaining discretion.'),
('Route and timing','Selects a conditional route and independently tracks each clock.'),
('Bounded request','Asks the proper recipient for an act within their assigned power.'),
('Remedy verification','Defines a result that can be checked and identifies downstream effects.'),
('Participation and revision','Protects privacy, considers practical effects, accepts challenge, and revises when evidence changes.')
],[1.65,5.39],9.8)
note('Score each dimension 0 to 4: 0 absent or wrong; 1 mostly unsupported; 2 partly correct with significant gaps; 3 accurate and usable with a stated gap; 4 accurate, sourced, and complete for the packet. Target 18 of 24, with at least 3 on authority and on bounded relief. Completion demonstrates practice of a skill, not legal qualification.')

# Sources appended separately to allow verified agent register updates
SOURCES = [
('S1','Founding purpose','Declaration of Independence overview','National Archives. Political foundation and explanation of consent and rights; the Declaration is not a legally binding case-specific filing vehicle.','https://www.archives.gov/founding-docs/declaration'),
('S2','Federal constitutional structure','Constitution and amendment transcripts','National Archives. Articles I, II, III and VI, the Tenth Amendment, and the Fourteenth Amendment provide distinct structural and rights sources.','https://www.archives.gov/founding-docs/constitution-transcript\nhttps://www.archives.gov/founding-docs/bill-of-rights-transcript\nhttps://www.archives.gov/founding-docs/amendments-11-27'),
('S3','Florida constitutional structure','Constitution of the State of Florida','Florida Legislature. Article I and Articles II through V address rights and the state government structure. Check the exact provision relevant to the claimed power.','https://www.leg.state.fl.us/statutes/index.cfm?submenu=3'),
('S4','Federal agency process','5 USC 555','Government Publishing Office. Section 555(b) concerns concluding a matter presented to an agency within a reasonable time; other subsections concern related administrative procedure.','https://www.govinfo.gov/link/uscode/5/555'),
('S5','APA coverage and limits','5 USC 701 and 702','Government Publishing Office. Exceptions, agency definitions, review framework, and limitations must be read together; this is not a universal state-court remedy.','https://www.govinfo.gov/content/pkg/USCODE-2024-title5/html/USCODE-2024-title5-partI-chap7-sec701.htm\nhttps://www.govinfo.gov/content/pkg/USCODE-2024-title5/html/USCODE-2024-title5-partI-chap7-sec702.htm'),
('S6','APA reviewability','5 USC 704','Government Publishing Office. Reviewability, final agency action, another adequate remedy in a court, and related provisions. Special statutory review routes may matter.','https://www.govinfo.gov/content/pkg/USCODE-2024-title5/html/USCODE-2024-title5-partI-chap7-sec704.htm'),
('S7','APA action and omission','5 USC 706 and Norton v Southern Utah Wilderness Alliance 542 US 55 2004','Government Publishing Office and Library of Congress. Section 706(1) and Norton distinguish a discrete legally required act from broad programmatic supervision or an order compelling a discretionary outcome.','https://www.govinfo.gov/content/pkg/USCODE-2024-title5/html/USCODE-2024-title5-partI-chap7-sec706.htm\nhttps://tile.loc.gov/storage-services/service/ll/usrep/usrep542/usrep542055/usrep542055.pdf'),
('S8','Florida dependency framework','Chapter 39 Proceedings Relating to Children','Florida Legislature, 2026 statutes. See 39.013 for jurisdiction; 39.401 for custody; 39.501 for petition; 39.507 for adjudication; and 39.521 for disposition. The checkpoint table names relevant subsections.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0000-0099/0039/0039.html'),
('S9','Florida shelter proceedings','Florida Statutes section 39.402','Florida Legislature. Contains shelter-related conditions and procedures. Do not transplant one stage’s standard into another.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0000-0099/0039/Sections/0039.402.html'),
('S10','Florida judicial review','Florida Statutes section 39.701','Florida Legislature, 2026 statutes. Review reports, determinations, and orders. The worked example uses the conditional duties in subsection (2)(d)3; return-home conditions remain separately applicable.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0000-0099/0039/Sections/0039.701.html'),
('S11','Florida dependency rules','Florida Rules of Juvenile Procedure effective October 1 2026','The Florida Bar official compilation. Rules 8.260, 8.265, 8.270, and 8.415 support the written-order, rehearing, relief, and judicial-review discussions. Rule 8.265(b)(3) says rehearing does not toll appeal time.','https://www-media.floridabar.org/uploads/2026/09/2027_04-OCT-Florida-Rules-of-Juvenile-Procedure-10-1-2026.pdf'),
('S12','Florida review and appeal','Florida Statutes section 39.510 and Florida Rules of Appellate Procedure','Florida Legislature and The Florida Bar. Distinguish the statute authorizing review from rules governing the vehicle, order, rendition, deadlines, and record.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0000-0099/0039/Sections/0039.510.html\nhttps://www-media.floridabar.org/uploads/2026/10/Appellate-Court-Rules-10-01-26.pdf'),
('S13','Florida dependency confidentiality','Florida Statutes section 39.0132','Florida Legislature. Contains confidentiality and records provisions. Verify access and disclosure rules before sharing actual case material.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0000-0099/0039/Sections/0039.0132.html'),
('S14','Florida court geography','Second District Court of Appeal history and jurisdiction','Florida Second District Court of Appeal. Its listed jurisdiction includes Hillsborough County and the Thirteenth Judicial Circuit. Appealability is a separate question.','https://2dca.flcourts.gov/About-the-Court/History-of-the-Court'),
('S15','Hillsborough trial court','Thirteenth Judicial Circuit of Florida','Official court site. Starting point for the court’s divisions, administrative orders, and local information. The relevant case and route still require verification.','https://www.fljud13.org/'),
('S16','Florida recording law','Florida Statutes section 934.03','Florida Legislature, 2026 statutes. Section 934.03(2)(d) includes an all-parties prior consent route. Review relevant definitions, other exceptions, and the nature of the communication before applying the law.','https://www.leg.state.fl.us/Statutes/index.cfm?App_mode=Display_Statute&URL=0900-0999/0934/Sections/0934.03.html')
]

for si,chunk in enumerate([SOURCES[:6],SOURCES[6:11],SOURCES[11:]]):
    page('Primary source register' + ('' if si==0 else ' continued'),'Reference '+str(si+1)+' of 3')
    if si==0:
        p('Source review date 9 October 2026. Sources support the bounded propositions described here. They are not a complete research record for an individual case. Recheck current text, effective dates, amendments, controlling decisions, and local rules before real-world use.',size=9.7)
    for sid,topic,title,desc,url in chunk:
        h(sid+' '+topic)
        p(title+'. '+desc,size=9.7)
        for u in url.split('\n'):
            url_paragraph(u)
    if si==2:
        h('Maintaining the course')
        p('Keep the website, fieldbook, slides, narration, and answer keys linked to the same source IDs. Record a change log when a source changes. Recheck affected predicate cards and exercises; identify the human reviewer and publication date. A verified link is evidence of a source location, not proof that every case falls within that source.',size=9.7)

doc.core_properties.title='Authority to Remedy Fieldbook'
doc.core_properties.subject='Volume 1 classroom fieldbook for tracing consequential public decisions'
doc.core_properties.author='Kapukai Governance Lab'
doc.core_properties.keywords='authority, duty, remedy, public decisions, classroom, blueprint'
doc.core_properties.comments=''
doc.save(OUT)
print(OUT)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
